# Builds ETB2-blade-profile.xlsx: the simple (linear chord and twist) NACA 0012 blade
# that students use as input to the CAD tutorial.
# Usage: python3 make_blade_xlsx.py ETB2-blade-profile.xlsx
import sys
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series

out = sys.argv[1] if len(sys.argv) > 1 else "ETB2-blade-profile.xlsx"
wb = Workbook()
B = Font(bold=True); H = Font(bold=True, color="FFFFFF"); T = Font(bold=True, size=14, color="004896")
INP = PatternFill("solid", fgColor="FFF2CC")      # inputs: yellow
HDR = PatternFill("solid", fgColor="004896")
NOTE = Font(italic=True, color="666666")
thin = Side(style="thin", color="BBBBBB"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)

def header(ws, row, labels, col=1):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=col+i, value=t); c.font = H; c.fill = HDR
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

# ---------------------------------------------------------------- Blade
ws = wb.active; ws.title = "Blade"
ws["A1"] = "ETB2 — simple blade profile (input for the CAD tutorial)"; ws["A1"].font = T
ws["A2"] = "Yellow cells are inputs — change them and the table below updates. Everything else is a formula."
ws["A2"].font = NOTE

header(ws, 3, ["Input", "Value", "Units"])
# (label, default, units, number format); row = 4 + index
inputs = [
    ("Tip radius R_tip",                 100,       "mm",    "0.0"),     # row 4
    ("Root station r_root (aerofoil starts)", 20,   "mm",    "0.0"),     # row 5
    ("Hub diameter",                     36,        "mm",    "0.0"),     # row 6
    ("Hub length",                       20,        "mm",    "0.0"),     # row 7
    ("Number of blades",                 3,         "",      "0"),       # row 8
    ("Root chord c_root",                24,        "mm",    "0.0"),     # row 9
    ("Tip chord c_tip",                  9,         "mm",    "0.0"),     # row 10
    ("Root twist",                       25,        "deg",   "0.0"),     # row 11
    ("Tip twist",                        5,         "deg",   "0.0"),     # row 12
    ("Aerofoil section",                 "NACA0012", "",     "@"),       # row 13
    ("Pitch axis (fraction of chord)",   0.25,      "c",     "0.00"),    # row 14
    ("Number of stations N",             9,         "",      "0"),       # row 15
]
for k, (lab, val, unit, fmt) in enumerate(inputs):
    r = 4 + k
    ws.cell(row=r, column=1, value=lab).border = BOX
    v = ws.cell(row=r, column=2, value=val); v.fill = INP; v.number_format = fmt; v.border = BOX
    ws.cell(row=r, column=3, value=unit).border = BOX

# Absolute references to the input cells
R_TIP, R_ROOT = "$B$4", "$B$5"
C_ROOT, C_TIP = "$B$9", "$B$10"
TW_ROOT, TW_TIP = "$B$11", "$B$12"
SECTION, N = "$B$13", "$B$15"

# Station table (fixed 9 rows: the row count does not follow N; add or delete rows to change N)
HROW = 18
ws.cell(row=HROW-1, column=1, value="Stations (all formulas; the table has 9 rows, so if you change N, add or delete rows by copying the last row).").font = NOTE
header(ws, HROW, ["Station i", "r (mm)", "r/R", "chord c (mm)", "twist (deg)", "section"])
NROWS = 9
first, last = HROW + 1, HROW + NROWS
for k in range(NROWS):
    r = first + k
    ws.cell(row=r, column=1, value=f"=ROW()-ROW($A${HROW})").number_format = "0"
    ws.cell(row=r, column=2, value=f"={R_ROOT}+({R_TIP}-{R_ROOT})*(A{r}-1)/({N}-1)").number_format = "0.000"
    ws.cell(row=r, column=3, value=f"=B{r}/{R_TIP}").number_format = "0.000"
    ws.cell(row=r, column=4, value=f"={C_ROOT}+({C_TIP}-{C_ROOT})*(B{r}-{R_ROOT})/({R_TIP}-{R_ROOT})").number_format = "0.000"
    ws.cell(row=r, column=5, value=f"={TW_ROOT}+({TW_TIP}-{TW_ROOT})*(B{r}-{R_ROOT})/({R_TIP}-{R_ROOT})").number_format = "0.00"
    ws.cell(row=r, column=6, value=f"={SECTION}").number_format = "@"
    for c in range(1, 7):
        ws.cell(row=r, column=c).border = BOX

for col, w in zip("ABCDEF", [36, 12, 10, 14, 12, 12]):
    ws.column_dimensions[col].width = w
ws.row_dimensions[HROW].height = 30

# Chart: chord and twist against r
ch = ScatterChart(); ch.title = "Chord and twist vs radius"; ch.style = 13
ch.x_axis.title = "r (mm)"; ch.y_axis.title = "chord (mm) / twist (deg)"
ch.x_axis.delete = False; ch.y_axis.delete = False
xs = Reference(ws, min_col=2, min_row=first, max_row=last)
for col, name in ((4, "chord c (mm)"), (5, "twist (deg)")):
    s = Series(Reference(ws, min_col=col, min_row=first, max_row=last), xs, title=name)
    s.marker.symbol = "circle"; s.smooth = False
    ch.series.append(s)
ch.height = 9; ch.width = 16
ws.add_chart(ch, "H3")

# ---------------------------------------------------------------- NACA 0012 section
wn = wb.create_sheet("NACA0012 section")
wn["A1"] = "NACA 0012 section coordinates"; wn["A1"].font = T
wn["A2"] = ("Normalised coordinates (unit chord): x/c runs from 0 at the leading edge to 1 at the trailing edge. "
            "Upper surface LE to TE, then lower surface TE to LE. Import into CAD or read in code.")
wn["A2"].font = NOTE
header(wn, 4, ["x/c", "y/c"])
t = 0.12   # 12% thick -> NACA 0012; chord = 1
xs_ = [0.5 * (1 - math.cos(math.pi * i / 40)) for i in range(41)]   # cosine spacing
half = [5 * t * (0.2969 * math.sqrt(x) - 0.1260 * x - 0.3516 * x**2 + 0.2843 * x**3 - 0.1015 * x**4)
        for x in xs_]
upper = list(zip(xs_, half))                                    # LE -> TE, +y
lower = [(x, -h) for x, h in reversed(upper[1:])]              # TE -> LE, -y (LE point not repeated)
for r, (x, y) in enumerate(upper + lower, start=5):
    wn.cell(row=r, column=1, value=x).number_format = "0.00000"
    wn.cell(row=r, column=2, value=y).number_format = "0.00000"
wn.column_dimensions["A"].width = 12; wn.column_dimensions["B"].width = 12

wb.save(out); print("saved", out)
